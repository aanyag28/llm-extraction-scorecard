from openpyxl import load_workbook
from pathlib import Path

BASE_DIR = Path.home() / "Documents" / "LLM-Extraction-Scorecard"
ANSWER_KEY = BASE_DIR / "answer_key" / "answer_key.csv-2.xlsx"

workbook = load_workbook(ANSWER_KEY)
sheet = workbook["Answer keys"]

headers = {}
for cell in sheet[1]:
    if cell.value is not None:
        headers[str(cell.value).strip()] = cell.column

filing_col = headers["Filing"]
ticker_col = headers["Ticker"]
metric_col = headers["Metric"]
value_col = headers["Value"]
unit_col = headers["Unit"]

production_values = {
    ("2025 10-K", "FANG"): 336178,
    ("Q1 2026 10-Q", "FANG"): 88142,
    ("Q2 2026 10-Q", "FANG"): 92607,

    ("2025 10-K", "PR"): 143311,
    ("Q1 2026 10-Q", "PR"): 37156,
    ("Q2 2026 10-Q", "PR"): 34253,

    ("2025 10-K", "MTDR"): 75581,
    ("Q3 2025 10-Q", "MTDR"): 19245,
    ("Q1 2026 10-Q", "MTDR"): 18683,

    ("2025 10-K", "CTRA"): 285576,
    ("Q3 2025 10-Q", "CTRA"): 210800,
    ("Q1 2026 10-Q", "CTRA"): 69400,

    ("2025 10-K", "SM"): 75500,
    ("Q1 2026 10-Q", "SM"): 33400,
    ("Q2 2026 10-Q", "SM"): 40000,
}

changed = 0

for row in range(2, sheet.max_row + 1):
    filing = sheet.cell(row, filing_col).value
    ticker = sheet.cell(row, ticker_col).value
    metric = sheet.cell(row, metric_col).value

    if not filing or not ticker or not metric:
        continue

    if str(metric).strip().lower() != "production":
        continue

    key = (str(filing).strip(), str(ticker).strip().upper())

    if key in production_values:
        sheet.cell(row, value_col).value = production_values[key]
        sheet.cell(row, unit_col).value = "MBOE"

        changed += 1

        print(
            f"Fixed: {filing} | {ticker} | "
            f"{production_values[key]} MBOE"
        )

workbook.save(ANSWER_KEY)

print()
print("====================================")
print("ANSWER KEY UPDATED")
print("====================================")
print(f"Production rows fixed: {changed}")
print()
print("Saved to:")
print(ANSWER_KEY)
