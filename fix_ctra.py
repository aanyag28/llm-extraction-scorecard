from openpyxl import load_workbook
from pathlib import Path

path = Path.home() / "Documents" / "LLM-Extraction-Scorecard" / "answer_key" / "answer_key.csv-2.xlsx"

workbook = load_workbook(path)
sheet = workbook["Answer keys"]

for row in range(2, sheet.max_row + 1):
    ticker = sheet.cell(row, 2).value
    filing = sheet.cell(row, 3).value
    metric = sheet.cell(row, 4).value

    if (
        ticker == "CTRA"
        and str(filing).strip() == "Q3 2025 10-Q"
        and metric == "Production"
    ):
        sheet.cell(row, 5).value = 72200
        sheet.cell(row, 6).value = "MBOE"
        print("Fixed CTRA Q3 2025 Production to 72,200 MBOE")

workbook.save(path)

print("Saved.")
