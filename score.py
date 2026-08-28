import csv
import re
from pathlib import Path
from openpyxl import load_workbook

# ============================================================
# SETUP
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
ANSWER_KEY = BASE_DIR / "answer_key" / "answer_key.csv-2.xlsx"
RESULTS_DIR = BASE_DIR / "results"
OUTPUT_FILE = RESULTS_DIR / "detailed_results.csv"

RESULTS_DIR.mkdir(exist_ok=True)

# ============================================================
# LOAD ANSWER KEY
# ============================================================

workbook = load_workbook(
    ANSWER_KEY,
    data_only=True
)

sheet = workbook["Answer keys"]

headers = [cell.value for cell in sheet[1]]

answer_rows = []

for row in sheet.iter_rows(min_row=2, values_only=True):
    data = dict(zip(headers, row))

    if data.get("Filing") and data.get("Metric"):
        answer_rows.append(data)


# ============================================================
# FIND GEMINI RESULT FILE
# ============================================================

def find_gemini_file(ticker, filing):

    ticker = str(ticker).strip().upper()
    filing = str(filing).strip().upper()

    if "2025 10-K" in filing:
        filename = f"{ticker}_10K_2025_result.txt"

    elif "Q1 2026 10-Q" in filing:
        filename = f"{ticker}_10Q_2026Q1_result.txt"

    elif "Q2 2026 10-Q" in filing:
        filename = f"{ticker}_10Q_2026Q2_result.txt"

    elif "Q3 2025 10-Q" in filing:
        filename = f"{ticker}_10Q_2025Q3_result.txt"

    else:
        return None

    path = RESULTS_DIR / filename

    if path.exists():
        return path

    return None


# ============================================================
# FIND METRIC SECTION
# ============================================================

def get_metric_section(text, metric):

    metric = metric.lower().strip()

    if metric == "revenue":

        number = "1"

        names = [
            r"revenue",
            r"revenues",
            r"total revenue",
            r"total revenues",
            r"total operating revenues",
            r"total operating revenues and other income",
            r"total revenues and other income"
        ]

    elif metric == "production":

        number = "2"

        names = [
            r"production",
            r"total production",
            r"total production volume",
            r"total oil equivalent",
            r"oil equivalent"
        ]

    elif metric == "capex":

        number = "3"

        names = [
            r"capex",
            r"capital expenditures",
            r"capital expenditure",
            r"total capital expenditures",
            r"capital expenditure activities"
        ]

    else:
        return ""

    name_pattern = "|".join(names)

    # --------------------------------------------------------
    # Format 1
    # ### Metric 1: Revenue
    # --------------------------------------------------------

    pattern1 = (
        rf"###\s*\**\s*Metric\s*{number}\s*:\s*"
        rf"(?:{name_pattern}).*?"
        rf"(?=###\s*\**\s*(?:Metric\s*[123]|\d+\s*[\.\:])|\Z)"
    )

    match = re.search(
        pattern1,
        text,
        flags=re.IGNORECASE | re.DOTALL
    )

    if match:
        return match.group(0)

    # --------------------------------------------------------
    # Format 2
    # ### 1. REVENUE
    # --------------------------------------------------------

    pattern2 = (
        rf"###\s*\**\s*{number}\s*[\.\:]\s*"
        rf"(?:{name_pattern}).*?"
        rf"(?=###\s*\**\s*\d+\s*[\.\:]|\Z)"
    )

    match = re.search(
        pattern2,
        text,
        flags=re.IGNORECASE | re.DOTALL
    )

    if match:
        return match.group(0)

    # --------------------------------------------------------
    # Format 3
    # METRIC: REVENUE
    # --------------------------------------------------------

    pattern3 = (
        rf"METRIC\s*:\s*(?:{name_pattern}).*?"
        rf"(?=METRIC\s*:|\Z)"
    )

    match = re.search(
        pattern3,
        text,
        flags=re.IGNORECASE | re.DOTALL
    )

    if match:
        return match.group(0)

    # --------------------------------------------------------
    # Format 4
    # Metric: Revenue
    # --------------------------------------------------------

    pattern4 = (
        rf"Metric\s*:\s*(?:{name_pattern}).*?"
        rf"(?=Metric\s*:|\Z)"
    )

    match = re.search(
        pattern4,
        text,
        flags=re.IGNORECASE | re.DOTALL
    )

    if match:
        return match.group(0)

    return ""


# ============================================================
# EXTRACT EXACT AMOUNT
# ============================================================

def extract_exact_amount(section):

    if not section:
        return "Not found"

    lines = section.splitlines()

    for i, line in enumerate(lines):

        cleaned = (
            line
            .replace("**", "")
            .replace("*", "")
            .replace("`", "")
            .strip()
        )

        match = re.search(
            r"Exact\s+Amount\s*:\s*(.*)",
            cleaned,
            re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            if value:
                return value

            # Value may be on next line
            for next_line in lines[i + 1:i + 5]:

                next_cleaned = (
                    next_line
                    .replace("**", "")
                    .replace("*", "")
                    .replace("`", "")
                    .strip()
                )

                if re.search(r"\d", next_cleaned):
                    return next_cleaned

    return "Not found"


# ============================================================
# EXTRACT UNIT
# ============================================================

def extract_unit(section):

    if not section:
        return ""

    lines = section.splitlines()

    for line in lines:

        cleaned = (
            line
            .replace("**", "")
            .replace("*", "")
            .replace("`", "")
            .strip()
        )

        match = re.search(
            r"Unit\s*:\s*(.+)",
            cleaned,
            re.IGNORECASE
        )

        if match:
            return match.group(1).strip()

    return ""


# ============================================================
# EXTRACT NUMBER
# ============================================================

def extract_number(value):

    if value is None:
        return None

    text = str(value)

    text = text.replace(",", "")
    text = text.replace("$", "")

    # Handle values such as:
    # TOTAL PRODUCTION = 285576 MBOE

    match = re.search(
        r"=\s*(-?\d+(?:\.\d+)?)",
        text
    )

    if match:
        return float(match.group(1))

    # Otherwise take the first number

    match = re.search(
        r"-?\d+(?:\.\d+)?",
        text
    )

    if not match:
        return None

    return float(match.group(0))


# ============================================================
# DETECT MONEY MULTIPLIER
# ============================================================

def detect_money_multiplier(value, unit):

    combined = f"{value} {unit}".lower()

    # Billion
    if "billion" in combined:
        return 1_000_000_000

    # Million
    if "million" in combined:
        return 1_000_000

    # Thousand
    if "thousand" in combined:
        return 1_000

    # Explicit dollars / USD with no stated larger unit
    if (
        "usd" in combined
        or "dollar" in combined
        or "$" in str(value)
    ):
        return 1

    return None


# ============================================================
# DETECT PRODUCTION MULTIPLIER
# ============================================================

def detect_production_multiplier(value, unit):

    combined = f"{value} {unit}".lower()

    combined = combined.replace(" ", "")

    # Daily production is NOT total production
    if "mboe/d" in combined:
        return None

    if "boe/d" in combined:
        return None

    # MMBOE = 1,000 MBOE
    if "mmboe" in combined:
        return 1000.0

    if "mmboe" in combined:
        return 1000.0

    # MBOE
    if "mboe" in combined:
        return 1.0

    # MBoe
    if "mboe" in combined:
        return 1.0

    return None


# ============================================================
# NORMALIZE GEMINI VALUE
# ============================================================

def normalize_value(value, unit, metric):

    if value is None:
        return None

    if str(value).strip().lower() in [
        "",
        "not found",
        "not reported",
        "not explicitly reported"
    ]:
        return None

    number = extract_number(value)

    if number is None:
        return None

    metric = metric.lower().strip()

    # --------------------------------------------------------
    # PRODUCTION
    # --------------------------------------------------------

    if metric == "production":

        multiplier = detect_production_multiplier(
            value,
            unit
        )

        if multiplier is None:
            return None

        # Convert to MBOE
        return number * multiplier

    # --------------------------------------------------------
    # MONEY
    #
    # Everything is converted to actual dollars.
    # --------------------------------------------------------

    multiplier = detect_money_multiplier(
        value,
        unit
    )

    if multiplier is None:
        return None

    return number * multiplier


# ============================================================
# NORMALIZE ANSWER KEY VALUE
# ============================================================

def normalize_answer_key(value, unit, metric):

    if value is None:
        return None

    number = extract_number(value)

    if number is None:
        return None

    metric = metric.lower().strip()

    # --------------------------------------------------------
    # PRODUCTION
    # --------------------------------------------------------

    if metric == "production":

        unit_text = str(unit).lower()

        if "mmboe" in unit_text:
            return number * 1000

        if "mboe" in unit_text:
            return number

        return None

    # --------------------------------------------------------
    # MONEY
    #
    # Convert answer key to actual dollars.
    # --------------------------------------------------------

    unit_text = str(unit).lower()

    if "billion" in unit_text:
        return number * 1_000_000_000

    if "million" in unit_text:
        return number * 1_000_000

    if "thousand" in unit_text:
        return number * 1_000

    if (
        "usd" in unit_text
        or "dollar" in unit_text
    ):
        return number

    return None


# ============================================================
# SCORE ONE ROW
# ============================================================

def score_row(
    extracted_value,
    extracted_unit,
    answer_value,
    answer_unit,
    metric
):

    extracted = normalize_value(
        extracted_value,
        extracted_unit,
        metric
    )

    answer = normalize_answer_key(
        answer_value,
        answer_unit,
        metric
    )

    if extracted is None:
        return False

    if answer is None:
        return False

    # --------------------------------------------------------
    # PRODUCTION
    # --------------------------------------------------------

    if metric.lower() == "production":

        return abs(extracted - answer) <= 1

    # --------------------------------------------------------
    # MONEY
    #
    # Allow 0.05% rounding tolerance.
    # --------------------------------------------------------

    tolerance = max(
        1.0,
        abs(answer) * 0.0005
    )

    return abs(extracted - answer) <= tolerance


# ============================================================
# CREATE DETAILED RESULTS
# ============================================================

detailed_rows = []

for answer in answer_rows:

    ticker = str(
        answer["Ticker"]
    ).strip()

    filing = str(
        answer["Filing"]
    ).strip()

    metric = str(
        answer["Metric"]
    ).strip()

    answer_value = answer["Value"]
    answer_unit = answer["Unit"]

    gemini_file = find_gemini_file(
        ticker,
        filing
    )

    if gemini_file is None:

        extracted_value = "Gemini result not found"
        extracted_unit = ""
        correct = False

    else:

        text = gemini_file.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        section = get_metric_section(
            text,
            metric
        )

        extracted_value = extract_exact_amount(
            section
        )

        extracted_unit = extract_unit(
            section
        )

        correct = score_row(
            extracted_value,
            extracted_unit,
            answer_value,
            answer_unit,
            metric
        )

    detailed_rows.append({

        "Filing": filing,

        "Ticker": ticker,

        "Metric": metric,

        "Gemini Extracted Value":
            extracted_value,

        "Gemini Unit":
            extracted_unit,

        "Answer Key Value":
            answer_value,

        "Answer Key Unit":
            answer_unit,

        "Correct":
            "Correct"
            if correct
            else "Incorrect"
    })


# ============================================================
# SAVE DETAILED RESULTS
# ============================================================

with open(
    OUTPUT_FILE,
    "w",
    newline="",
    encoding="utf-8"
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Filing",
            "Ticker",
            "Metric",
            "Gemini Extracted Value",
            "Gemini Unit",
            "Answer Key Value",
            "Answer Key Unit",
            "Correct"
        ]
    )

    writer.writeheader()
    writer.writerows(
        detailed_rows
    )


# ============================================================
# SCORE SUMMARY
# ============================================================

def rows_for(metric):

    return [
        row
        for row in detailed_rows
        if row["Metric"].lower().strip()
        == metric.lower().strip()
    ]


def count_correct(rows):

    return sum(
        row["Correct"] == "Correct"
        for row in rows
    )


def pct(rows):

    if not rows:
        return 0

    return (
        count_correct(rows)
        / len(rows)
        * 100
    )


revenue_rows = rows_for("Revenue")
production_rows = rows_for("Production")
capex_rows = rows_for("CapEx")

total = len(detailed_rows)

correct_total = sum(
    row["Correct"] == "Correct"
    for row in detailed_rows
)


# ============================================================
# PRINT SCORE
# ============================================================

print()

print("====================================")
print("GEMINI EXTRACTION SCORE")
print("====================================")

print(
    f"Revenue:     "
    f"{count_correct(revenue_rows)}/"
    f"{len(revenue_rows)} "
    f"({pct(revenue_rows):.1f}%)"
)

print(
    f"Production:  "
    f"{count_correct(production_rows)}/"
    f"{len(production_rows)} "
    f"({pct(production_rows):.1f}%)"
)

print(
    f"CapEx:       "
    f"{count_correct(capex_rows)}/"
    f"{len(capex_rows)} "
    f"({pct(capex_rows):.1f}%)"
)

if total:

    print(
        f"Overall:     "
        f"{correct_total}/"
        f"{total} "
        f"({correct_total / total * 100:.1f}%)"
    )

else:

    print(
        "Overall:     0/0 (0.0%)"
    )

print("====================================")

print()

print("Detailed results saved to:")
print(OUTPUT_FILE)