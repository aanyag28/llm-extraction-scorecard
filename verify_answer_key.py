from pathlib import Path
from openpyxl import load_workbook
import re

BASE_DIR = Path.home() / "Documents" / "LLM-Extraction-Scorecard"
ANSWER_KEY = BASE_DIR / "answer_key" / "answer_key.csv-2.xlsx"

workbook = load_workbook(ANSWER_KEY)
sheet = workbook["Answer keys"]

print()
print("====================================")
print("ANSWER KEY VERIFICATION")
print("====================================")

for row in range(2, sheet.max_row + 1):

    ticker = sheet.cell(row, 2).value
    filing = sheet.cell(row, 3).value
    metric = sheet.cell(row, 4).value
    value = sheet.cell(row, 5).value
    unit = sheet.cell(row, 6).value

    if not ticker or not filing or not metric:
        continue

    # Convert the answer-key filing name into the filename
    filing_clean = str(filing).strip()

    if filing_clean == "2025 10-K":
        filename = f"{ticker}_10K_2025.txt"

    elif filing_clean == "Q1 2026 10-Q":
        filename = f"{ticker}_10Q_2026Q1.txt"

    elif filing_clean == "Q2 2026 10-Q":
        filename = f"{ticker}_10Q_2026Q2.txt"

    elif filing_clean == "Q3 2025 10-Q":
        filename = f"{ticker}_10Q_2025Q3.txt"

    else:
        print(f"UNKNOWN FILING: {ticker} {filing}")
        continue

    filing_path = BASE_DIR / filename

    if not filing_path.exists():
        print()
        print(f"FILE NOT FOUND: {filing_path}")
        continue

    text = filing_path.read_text(
        encoding="utf-8",
        errors="ignore"
    )

    print()
    print("------------------------------------")
    print(f"{ticker} | {filing_clean} | {metric}")
    print(f"Answer key: {value} {unit}")
    print(f"File: {filename}")
    print("------------------------------------")

    # Production was already verified separately.
    if metric == "Production":
        print("Production: VERIFIED AS TOTAL MBOE")
        continue

    # Search for the numeric answer-key value in the filing.
    value_text = str(value)

    # Handle decimal values such as 69.4
    search_terms = [
        value_text,
        f"{value:,.0f}" if isinstance(value, (int, float)) else value_text,
        f"{value:,.1f}" if isinstance(value, float) else value_text
    ]

    found = False

    for term in search_terms:

        if term in text:
            index = text.find(term)

            start = max(0, index - 500)
            end = min(len(text), index + 800)

            print()
            print(f"FOUND VALUE: {term}")
            print()
            print(text[start:end])

            found = True
            break

    if not found:

        print()
        print("VALUE NOT FOUND EXACTLY IN FILING.")
        print()
        print("Searching for metric keywords...")

        if metric == "Revenue":
            keywords = [
                "total revenues",
                "total operating revenues",
                "total operating revenue",
                "oil and gas sales",
                "revenues"
            ]

        elif metric == "CapEx":
            keywords = [
                "capital expenditures",
                "capital expenditure",
                "drilling and development capital",
                "drilling, completion",
                "capital expenditures for drilling"
            ]

        else:
            keywords = []

        keyword_found = False

        for keyword in keywords:

            match = re.search(
                re.escape(keyword),
                text,
                re.IGNORECASE
            )

            if match:

                start = max(0, match.start() - 500)
                end = min(len(text), match.start() + 1000)

                print()
                print(f"FOUND KEYWORD: {keyword}")
                print()
                print(text[start:end])

                keyword_found = True
                break

        if not keyword_found:
            print("No relevant keyword found.")

print()
print("====================================")
print("VERIFICATION COMPLETE")
print("====================================")
